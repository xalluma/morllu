"""Genera Zimia_Establecimientos_Rubi_V1.xlsx y .csv a partir de los registros verificados.

Los registros se han consolidado a mano a partir de:
  - Overture Maps Places, release 2026-09-23.0 (coordenadas y datos de contacto agregados
    de Meta, Microsoft, Foursquare, AllThePlaces y PinMeTo), filtrado con el polígono
    municipal de Rubí (Overture divisions, division_area id 7f6eeace-47c8-41b9-bf66-82185f798ef8).
  - Búsquedas web del 2026-10-14 (webs oficiales, localizadores de cadenas, rubicomerc.cat y
    directorios). Ver notas_evidencias_2026-10-14.md.

Uso: python3 build_rubi_v1.py <directorio_salida>
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HOY = "2026-10-14"
OVT = "Overture Maps Places 2026-09-23.0"

CAMPOS = [
    "id_zimia", "nombre_comercial", "tipo_establecimiento", "tipo_empresa", "nombre_cadena",
    "direccion", "codigo_postal", "municipio", "provincia", "comunidad_autonoma", "pais",
    "latitud", "longitud", "telefono", "web", "email",
    "vende_alimentacion_mascotas", "perros", "gatos", "otros_animales",
    "fuente_principal", "url_fuente", "fuentes_secundarias", "fecha_obtencion",
    "fecha_ultima_verificacion", "estado_registro", "nivel_verificacion", "metodo_verificacion",
    "observaciones_internas", "fecha_actualizacion", "reclamado_por_comercio",
]
assert len(CAMPOS) == 31

PERMITIDOS = {
    "tipo_establecimiento": {"tienda_mascotas", "veterinario", "supermercado", "hipermercado",
                             "agrotienda", "garden", "tienda_natural", "otros"},
    "tipo_empresa": {"independiente", "cadena", None},
    "vende_alimentacion_mascotas": {"si", "no", "no_verificado"},
    "perros": {"si", "no", "no_verificado"},
    "gatos": {"si", "no", "no_verificado"},
    "otros_animales": {"si", "no", "no_verificado"},
    "estado_registro": {"activo", "cerrado", "pendiente_revision"},
    "nivel_verificacion": {"confirmado", "parcialmente_confirmado", "no_verificado"},
    "metodo_verificacion": {"fuente oficial", "fuente externa", "comercio", "manual", None},
    "reclamado_por_comercio": {"false", "true"},
}

NV = "no_verificado"
CHAIN_FOOD_NOTE = ("Venta de alimentación para mascotas no verificada en esta sucursal: solo hay "
                   "indicio genérico de la categoría (supermercado de cadena). Confirmar en tienda o "
                   "con el surtido de la sucursal.")


def ovt(oid, origen):
    return f"{OVT} (registro '{oid}'; origen {origen})"


R = []


def rec(**kw):
    base = dict(
        tipo_empresa=None, nombre_cadena=None, codigo_postal="08191", municipio="Rubí",
        provincia="Barcelona", comunidad_autonoma="Cataluña", pais="España",
        latitud=None, longitud=None, telefono=None, web=None, email=None,
        vende_alimentacion_mascotas=NV, perros=NV, gatos=NV, otros_animales=NV,
        fuentes_secundarias=None, fecha_obtencion=HOY, fecha_ultima_verificacion=HOY,
        estado_registro="pendiente_revision", nivel_verificacion=NV,
        metodo_verificacion="fuente externa", observaciones_internas=None,
        fecha_actualizacion=HOY, reclamado_por_comercio="false",
    )
    base.update(kw)
    # Campos de trabajo (no forman parte de la estructura V1): hoja de pendientes.
    base.setdefault("_pend_motivo", None)
    base.setdefault("_pend_info", None)
    R.append(base)


# ---------------------------------------------------------------- TIENDAS DE MASCOTAS / AGRO
rec(nombre_comercial="Sweet Doggy", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer de Sant Jaume, 32", latitud=41.490922, longitud=2.032061,
    telefono="+34931799570", web="http://www.sweetdoggy.cat/",
    vende_alimentacion_mascotas="si", perros="si", gatos="si",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat), ficha publicada por el comercio",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=ca&id=950",
    fuentes_secundarias="; ".join([
        ovt("DoggyRade", "Meta; marca 'DoggyRade'"),
        "https://animalia.com.es/rubi/tiendas-productos-para-animales/sweet-doggy",
        "https://www.srperro.com/sweet-doggy/", "https://pelucanes.com/sweet-doggy-9062/",
        "https://firmania.es/rubi/sweet-doggy-1042692"]),
    estado_registro="activo", nivel_verificacion="confirmado",
    observaciones_internas=("Ficha del comercio: 'Alimentació natural per mascotes' y peluquería canina. "
                            "Reseñas de clientes citan pienso y snacks para perros y gatos. Algunos directorios lo "
                            "clasifican como peluquería. Overture lo etiqueta con la marca 'DoggyRade' (sin "
                            "confirmar). Web alternativa citada en un directorio: sweetdoggy.es."))

rec(nombre_comercial="Petslovers Rubí", tipo_establecimiento="tienda_mascotas", tipo_empresa=None,
    direccion="Carretera de Sabadell, km 12,7, Nave 9", latitud=41.50523, longitud=2.037931,
    telefono="+34633667555", web="https://petslovers.es/",
    vende_alimentacion_mascotas="si", perros="si", gatos="si", otros_animales="si",
    fuente_principal="Web oficial Petslovers", url_fuente="https://petslovers.es/",
    fuentes_secundarias="; ".join([
        "https://petslovers.es/blog/actualidad/8-aniversario-un-poquito-sobre-nosotros",
        ovt("Petslovers Rubi", "Meta, confianza 0,98"),
        "https://www.paginasamarillas.es/f/rubi/petslovers-rubi_223247354_000000019.html",
        "https://caninoteca.es/tiendas-de-animales/rubi/"]),
    estado_registro="activo", nivel_verificacion="confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Web oficial: tienda de pienso para perros y gatos; la empresa cita también pienso "
                            "para aves y roedores y su tienda de Rubí con parking propio. tipo_empresa NULL: no "
                            "se ha podido determinar si opera más tiendas físicas (ficha de Páginas Amarillas "
                            "asociada a 'Terranova CNC')."))

rec(nombre_comercial="Agrorosell", tipo_establecimiento="agrotienda", tipo_empresa="independiente",
    direccion="Carrer de Méndez Núñez, 10", latitud=41.488545, longitud=2.030616,
    telefono="+34936990135", web="http://agrorosell.com/",
    vende_alimentacion_mascotas="si", otros_animales="si",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=ca&id=211",
    fuentes_secundarias="; ".join([
        "https://www.iberinform.es/empresa/8236477/agrorosell-rubi",
        "https://www.qdq.com/agro-rosell-491389", "https://www.facebook.com/agrorosell/",
        "https://caninoteca.es/tiendas-de-animales/rubi/", ovt("Agrorosell", "Meta")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Agrorosell Rubí SL (B67340026, activa según Iberinform, dato 18/09/2025). Objeto "
                            "social: venta de piensos para animales de granja y de compañía. Directorios: "
                            "ganadería, agricultura, jardinería y pajarería. No hay mención explícita a pienso "
                            "de perro o gato; web tomada de la ficha de Meta, sin abrir."))

rec(nombre_comercial="Agrifauna", tipo_establecimiento="agrotienda", tipo_empresa="independiente",
    direccion="Carrer de Marconi, 3, bajo (local 1)", latitud=41.485847, longitud=2.036697,
    telefono="+34936971356", web="https://agrifauna.es/",
    vende_alimentacion_mascotas="si", otros_animales="si",
    fuente_principal="Web oficial Agrifauna", url_fuente="https://agrifauna.es/contacto/",
    fuentes_secundarias="; ".join([
        "https://www.paginasamarillas.es/f/rubi/agrifauna_011686813_000000001.html",
        "https://www.qdq.com/agri-fauna-1174191", "https://firmania.es/rubi/agrifauna-1394502",
        ovt("Agrifauna", "Microsoft, confianza 0,97")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Web oficial: 'Alimentación y productos para el cuidado de su mascota', plantel, "
                            "semillas e insecticidas; servicio a domicilio; L-V 9-13:30 y 17-20, S 9-14. "
                            "Pajarería (aves y exóticos) según directorios y reseñas. Falta confirmar pienso "
                            "de perro/gato."))

rec(nombre_comercial="Gosexpress", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Avinguda de Castellbisbal, 32", latitud=41.490444, longitud=2.025235,
    telefono="+34690323357", web="http://www.gosexpress.com/",
    vende_alimentacion_mascotas="si",
    fuente_principal="Página de Facebook del comercio", url_fuente="https://www.facebook.com/gosexpress/",
    fuentes_secundarias="; ".join([
        "https://mapcarta.com/N11346285191 (OpenStreetMap: shop=pet)",
        "https://muchamascota.es/barcelona/gosexpress-rubi/", "https://firmania.es/rubi/gosexpress-1227581",
        "https://caninoteca.es/tiendas-de-animales/rubi/", ovt("Gosexpress", "Meta")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("El comercio se define como 'especialistas en alimentación y accesorios para "
                            "mascota, con entrega a domicilio'. Teléfono discrepante: 690 32 33 57 (Meta y "
                            "varios directorios) frente a 620 47 44 77 (otros). Overture asigna el mismo "
                            "690 32 33 57 a Fauna L'Amistat: posible mismo titular. Falta detalle por especie."),
    _pend_motivo="Teléfono contradictorio entre fuentes", _pend_info="Confirmar teléfono y especies (perro/gato)")

rec(nombre_comercial="Fauna L'Amistat", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer de Lluís Ribas, 1", latitud=41.48804, longitud=2.03485,
    telefono="+34935886379", web="http://www.faunalamistat.cat",
    vende_alimentacion_mascotas="si",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=ca&id=434",
    fuentes_secundarias="; ".join([
        "https://www.qdq.com/fauna-la-amistad-599236",
        "https://www.espainfo.com/fauna-lamistat-F1108C1041DD642",
        "https://yellow.place/en/fauna-l-amistat-rub%C3%AD-spain", ovt("Fauna L'Amistat", "Meta")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("También 'Fauna La Amistad'. Directorios: pajarería, cría de canarios, venta de "
                            "animales, alimentación, acuarios y peluquería para perros y gatos. Discrepancias: "
                            "número de portal (1 en rubicomerc/QDQ, 3 en Overture; coordenadas de Overture) y "
                            "teléfono (Overture: 690 32 33 57, compartido con Gosexpress)."),
    _pend_motivo="Número de portal y teléfono contradictorios",
    _pend_info="Confirmar portal (1 o 3), teléfono y alimentación por especie")

rec(nombre_comercial="El Reino Animal", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer de Pere Esmendia, 10", latitud=41.489926, longitud=2.030891,
    telefono="+34936973313",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=es&id=485",
    fuentes_secundarias="; ".join([
        "https://www.qdq.com/el-reino-animal-318481", "https://www.cylex.es/rubi/el-reino-animal-13327530.html",
        "https://mascotashoy.es/el-reino-animal-rubi-barcelona/",
        "https://pininastyle.es/tienda-de-animales-en-rubi-el-reino-animal/", ovt("El Reino Animal", "Meta")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Tienda de animales con varias fichas coincidentes (4,6/5 con 78 valoraciones según "
                            "un directorio). La venta de comida solo aparece de forma genérica en directorios; "
                            "no se marca 'si'. No confundir con 'Veterinari El Regne Animal' (Ctra. de Sabadell "
                            "69, hoy Centre Veterinari Rubí Nord)."),
    _pend_motivo="Venta de alimentación no confirmada", _pend_info="Confirmar venta de pienso perro/gato")

rec(nombre_comercial="Boutique Animal", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Passeig de les Torres, 49, local 1", latitud=41.48107, longitud=2.03372,
    telefono="+34936990947",
    vende_alimentacion_mascotas="si",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=es&id=937",
    fuentes_secundarias="; ".join([
        "https://www.facebook.com/p/Boutique-Animal-100054557703926/",
        "https://www.paginasamarillas.es/f/rubi/boutique-animal_220955637_000000001.html",
        "https://www.cylex.es/rubi/boutique-animal-12864711.html",
        ovt("Tienda especializada en piensos complementos para mascotas y tienda hípica", "Meta")]),
    estado_registro="pendiente_revision", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Ficha del comercio con marcas de pienso (Royal Canin, Purina, Advance, Dingo...). "
                            "Su Facebook anuncia que la tienda está 'temporalmente cerrada por venta de local' "
                            "(aviso sin fecha visible). Overture lo recoge sin nombre comercial en la misma "
                            "dirección y teléfono. Otros teléfonos: 650 14 53 92."),
    _pend_motivo="Posible cierre (aviso en Facebook, sin fecha)",
    _pend_info="Confirmar si sigue abierta; perro/gato")

rec(nombre_comercial="Tot Pels Teus Amics", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Passeig de la Riera, 35", latitud=41.489535, longitud=2.027686,
    telefono="+34672490078", vende_alimentacion_mascotas="si",
    fuente_principal="Página de Facebook del comercio",
    url_fuente="https://www.facebook.com/p/Tot-Pels-Teus-Amics-100063616141816/",
    fuentes_secundarias="; ".join([
        "https://es.foursquare.com/v/tot-pels-teus-amics/4f4ab3f9e4b0b1e3522e5564",
        "https://www.rubicomerc.cat/comercos.php?lang=ca&id=606",
        "https://es.nextdoor.com/pages/tot-pels-teus-amics/", ovt("Tot Pels Teus Amics", "Meta")]),
    estado_registro="pendiente_revision", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Clasificación contradictoria: Facebook (pienso y animales), Foursquare y Nextdoor "
                            "(artículos para mascotas), rubicomerc (flores, plantas y abonos). Teléfono: Meta "
                            "672 49 00 78; otro directorio 934 08 97 09."),
    _pend_motivo="Actividad y teléfono contradictorios",
    _pend_info="Confirmar tipo de comercio, teléfono y venta de pienso perro/gato")

rec(nombre_comercial="El Flamenco Rosa", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Plaça de Josep Sapés i Esmendia, 2, local 2",
    telefono="+34935861593", email="flamenco.pets@gmail.com",
    vende_alimentacion_mascotas="si",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=es&id=808",
    fuentes_secundarias="; ".join([
        "https://www.facebook.com/El.Flamenco.Rosa/", "https://dogwell.es/el-flamenco-rosa/",
        "http://miperroesunico.com/peluquerias-caninas-para-perros/el-flamenco-rosa-peluqueria-canina-en-rubi.html"]),
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Ficha del comercio: pinsos, grans, cosmètica, collars; perruqueria canina; botiga "
                            "hípica; marcas Royal Canin, Advance, Arion. Email tomado de directorio. No aparece "
                            "en Overture 2026-09 (en el n.º 2 de la plaza figura una inmobiliaria): la "
                            "actividad actual no está confirmada. Sin coordenadas fiables."),
    _pend_motivo="Actividad actual no confirmada; sin coordenadas",
    _pend_info="Confirmar existencia, local y coordenadas")

rec(nombre_comercial="Rubí Fauna", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer de Miguel de Cervantes, 84", telefono="+34931665052", web="http://rubifauna.com",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=ca&id=796",
    fuentes_secundarias="https://www.guiaveterinarios.com/pro/rubi-fauna-rubi-3980479/",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Directorios: accesorios, alimentación y peluquería de perros y gatos (evidencia "
                            "antigua, sin fecha). No aparece en Overture 2026-09 ni en búsquedas recientes. "
                            "Web citada en directorio, no comprobada."),
    _pend_motivo="Posiblemente cerrado; evidencia antigua", _pend_info="Confirmar existencia y coordenadas")

rec(nombre_comercial="Cuida tus mascotas", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer de Getafe, 2", telefono="+34931665522",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=es&id=944",
    fuentes_secundarias="https://empresite.eleconomista.es/Actividad/TIENDAS-ANIMALES-DOMESTICOS/localidad/RUBI-BARCELONA/",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("La ficha (de hace unos 11 años) cita pienso para perros y gatos, roedores, "
                            "peluquería y autolavado. Evidencia antigua: no se marca venta. No aparece en "
                            "Overture 2026-09."),
    _pend_motivo="Evidencia de hace ~11 años", _pend_info="Confirmar existencia, actividad y coordenadas")

rec(nombre_comercial="Animal Paradis", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer del Doctor Robert, 23",
    fuente_principal="QDQ", url_fuente="https://www.qdq.com/animal-paradis-362277",
    fuentes_secundarias="https://www.tiendasmascotas.net/animal-paradis-rubi/",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Solo en directorios. Teléfonos contradictorios (931 756 401 / 630 136 064). No "
                            "aparece en Overture 2026-09; en Doctor Robert 22-25 figuran otros comercios."),
    _pend_motivo="Solo directorios; posiblemente cerrado", _pend_info="Confirmar existencia, teléfono y coordenadas")

rec(nombre_comercial="Aquadisseny Aquaris", tipo_establecimiento="tienda_mascotas", tipo_empresa="independiente",
    direccion="Carrer de Monturiol, 8A", latitud=41.487381, longitud=2.034114,
    telefono="+34936971766", web="http://www.aquadissenyaquaris.com/",
    fuente_principal="Páginas Amarillas",
    url_fuente="https://www.paginasamarillas.es/f/rubi/aquadisseny-aquaris_011305463_000000001.html",
    fuentes_secundarias=ovt("Aquadisseny Aquaris", "Meta, confianza 0,95"),
    estado_registro="pendiente_revision", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Tienda de acuariofilia (peces tropicales). Directorios contradictorios sobre si "
                            "está abierta. Probable alimentación para peces, no verificada; perros y gatos poco "
                            "probables."),
    _pend_motivo="Estado contradictorio", _pend_info="Confirmar actividad y venta de alimento (peces)")

# Peluquerías / servicios con posible venta
rec(nombre_comercial="Peludets Nets", tipo_establecimiento="otros", tipo_empresa="independiente",
    direccion="Carretera de Terrassa, 28", latitud=41.497652, longitud=2.031389, telefono="+34631532154",
    vende_alimentacion_mascotas="si",
    fuente_principal="Página de Facebook del comercio", url_fuente="https://www.facebook.com/PeludetsNets/",
    fuentes_secundarias="; ".join([
        "https://www.paginasamarillas.es/f/rubi/peludets-nets_231467762_000000001.html",
        "https://pelucanes.com/cataluna/rubi/", ovt("Peludets Nets", "Meta, confianza 0,91")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Peluquería canina y felina que, según su página, vende también alimentación y "
                            "artículos de higiene. Falta confirmar la alimentación por especie."))

rec(nombre_comercial="Peluquería canina Gosset.net", tipo_establecimiento="otros", tipo_empresa="independiente",
    direccion="Avinguda de Catalunya, 10", telefono="+34660930036",
    fuente_principal="Serveis Actius", url_fuente="https://www.serveisactius.cat/es/valles/peluqueria-canina-rubi",
    fuentes_secundarias="https://caninoteca.es/peluquerias-caninas/rubi/peluqueria-canina-gossetnet/; https://perrolia.es/barcelona/peluquerias-caninas/peluqueria-canina-gosset-net-09064bafe8",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Peluquería canina con productos de cuidado de la piel. Sin evidencia de venta de "
                            "alimentación. No localizada en Overture: sin coordenadas."),
    _pend_motivo="Peluquería sin evidencia de venta de alimento", _pend_info="Preguntar si vende pienso; coordenadas")

rec(nombre_comercial="La Miri i Els Seus Amics", tipo_establecimiento="otros", tipo_empresa="independiente",
    direccion="Carrer del Torrent de l'Alba, 3", latitud=41.490192, longitud=2.034323, telefono="+34686155524",
    fuente_principal="peluqueriacanina.com.es",
    url_fuente="https://peluqueriacanina.com.es/rubi/la-miri-i-els-seus-amics-en-rubi/",
    fuentes_secundarias="https://infoperros.es/peluquerias-caninas/la-miri-i-els-seus-amics-rubi/; " + ovt("La Miri I Els Seus Amics", "Meta, confianza 0,08"),
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Peluquería canina (Overture la clasifica erróneamente como veterinario). Sin "
                            "evidencia de venta de alimentación. Confianza de Overture muy baja (0,08); "
                            "coordenadas por revisar."),
    _pend_motivo="Peluquería sin evidencia de venta; coordenadas de baja confianza",
    _pend_info="Preguntar si vende pienso; confirmar ubicación")

rec(nombre_comercial="Miss Peluts", tipo_establecimiento="otros", tipo_empresa=None,
    direccion="Carrer de Magí Ramentol, 13", latitud=41.4898, longitud=2.03396, telefono="+34931792349",
    fuente_principal=OVT, url_fuente="https://docs.overturemaps.org/guides/places/",
    fuentes_secundarias="Overture id 99f66be1-8ff8-426b-aea0-17aa1202ab87 (origen Meta, confianza 0,61)",
    fecha_ultima_verificacion=None, estado_registro="pendiente_revision", nivel_verificacion=NV,
    metodo_verificacion=None,
    observaciones_internas=("Servicio para animales según Overture/Meta; la búsqueda web no devolvió "
                            "resultados. Actividad y oferta desconocidas."),
    _pend_motivo="Una sola fuente", _pend_info="Confirmar existencia, actividad y venta de alimento")

rec(nombre_comercial="Noelia Martín Estilista Canina", tipo_establecimiento="otros", tipo_empresa="independiente",
    direccion="Carrer de la Pastora, 23", latitud=41.496837, longitud=2.040454, telefono="+34931448548",
    fuente_principal=OVT, url_fuente="https://docs.overturemaps.org/guides/places/",
    fuentes_secundarias="Overture id 6844f7c3-64db-4d8e-b010-779eec2424a8 (origen Meta, confianza 0,84)",
    fecha_ultima_verificacion=None, estado_registro="pendiente_revision", nivel_verificacion=NV,
    metodo_verificacion=None,
    observaciones_internas=("Nombre en la fuente: 'NoeliamartinEstilistacanina' (peluquería canina; Overture la "
                            "clasifica como veterinario). No contrastado en web."),
    _pend_motivo="Una sola fuente", _pend_info="Confirmar existencia y venta de alimento")

rec(nombre_comercial="Melenudets", tipo_establecimiento="otros", tipo_empresa="independiente",
    direccion="Carretera de Sabadell, 1", latitud=41.50471, longitud=2.03742, telefono="+34675421534",
    web="https://linktr.ee/melenudets",
    fuente_principal="Empresite (Melenudets SL)", url_fuente="https://empresite.eleconomista.es/MELENUDETS.html",
    fuentes_secundarias=ovt("Melenudets", "Meta, confianza 0,28"),
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Melenudets SL (B26857987), CNAE 9609 'otros servicios personales' (probable "
                            "peluquería). Overture la clasifica como tienda de mascotas con confianza baja."),
    _pend_motivo="Actividad dudosa", _pend_info="Confirmar si es tienda o peluquería y si vende pienso")

rec(nombre_comercial="Animalots Rubí", tipo_establecimiento="otros", tipo_empresa=None,
    direccion="Carrer de Maria Aurèlia Capmany, 42 (esquina C/ Bailèn)", latitud=41.500608, longitud=2.034875,
    telefono="+34697610366",
    fuente_principal=OVT, url_fuente="https://docs.overturemaps.org/guides/places/",
    fuentes_secundarias="Overture registro 'Animalots Rubi' (origen Meta, confianza 0,85); directorio local de Rubí (dirección con esquina C/ Bailèn)",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas="Servicio para animales sin descripción de actividad. Tipo de negocio desconocido.",
    _pend_motivo="Actividad desconocida", _pend_info="Confirmar actividad y venta de alimento")

# ---------------------------------------------------------------- VETERINARIOS
rec(nombre_comercial="Dispensari Veterinari del Vallès - Rubí 1 (Hospital)", tipo_establecimiento="veterinario",
    tipo_empresa="cadena", nombre_cadena="Dispensari Veterinari del Vallès",
    direccion="Carretera de Sant Cugat, 5", latitud=41.486915, longitud=2.03753, telefono="+34936991182",
    web="https://www.dispensariveterinari.com/",
    vende_alimentacion_mascotas="si", perros="si", gatos="si", otros_animales="si",
    fuente_principal="Web oficial Dispensari Veterinari del Vallès",
    url_fuente="https://www.dispensariveterinari.com/botigues/botiga-veterinaria-a-rubi/",
    fuentes_secundarias="; ".join([
        "https://www.dispensariveterinari.com/es/veterinarios/clinica-veterinaria-urgencias/",
        "https://www.rubicomerc.cat/comercos.php?lang=ca&id=590",
        ovt("Dispensari Veterinari del Vallès", "Meta, confianza 1,0")]),
    estado_registro="activo", nivel_verificacion="confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Hospital 24 h con tienda: pienso de gama alta para perros, gatos, conejos, "
                            "roedores, hurones, aves, peces y tortugas; entrega a domicilio. El grupo tiene "
                            "centros en Rubí, Sant Cugat y Cerdanyola."))

rec(nombre_comercial="Dispensari Veterinari del Vallès - Rubí 2", tipo_establecimiento="veterinario",
    tipo_empresa="cadena", nombre_cadena="Dispensari Veterinari del Vallès",
    direccion="Carrer del Gimnàs, 15", latitud=41.492334, longitud=2.02948, telefono="+34935882887",
    web="https://www.dispensariveterinari.com/veterinaris/rubi-2/",
    vende_alimentacion_mascotas="si",
    fuente_principal="Web oficial Dispensari Veterinari del Vallès",
    url_fuente="https://www.dispensariveterinari.com/botigues/botiga-veterinaria-a-rubi-2/",
    fuentes_secundarias="; ".join([
        "https://www.rubicomerc.cat/comercos.php?lang=ca&id=666",
        ovt("SonsoVet Homeopatia Acupuntura Veterinaria", "Meta; C/ Gimnàs 17, mismo teléfono")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Tienda con pienso medicado (seco y húmedo) y productos para mascotas, según la "
                            "web oficial; no detalla especies. Overture lo recoge como 'SonsoVet' en el n.º 17 "
                            "con el mismo teléfono (fusionado; coordenadas de ese registro)."))

rec(nombre_comercial="Centre Veterinari Les Torres", tipo_establecimiento="veterinario",
    tipo_empresa="cadena", nombre_cadena="IVC Evidensia",
    direccion="Avinguda de Can Cabanyes, 2", latitud=41.485796, longitud=2.031377, telefono="+34935883002",
    web="https://www.vetertorres.com/", email="recepcio@vetertorres.com",
    fuente_principal="Web oficial Centre Veterinari Les Torres", url_fuente="https://www.vetertorres.com/contacto/",
    fuentes_secundarias="; ".join([
        "https://evidensia.es/rubi/centre-veterinari-les-torres",
        "https://es.jooble.org/jdp/33677782309405124 (oferta IVC Evidensia 'CV Les Torres')",
        ovt("Centre Veterinari Les Torres", "PinMeTo, operating_status=open, confianza 1,0")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Urgencias 609 343 203. L-S 9-21 según la web (otro directorio: sábado hasta las "
                            "15). Sin evidencia de venta de alimentación. Overture tiene un duplicado 'Las Torres' "
                            "en Pg. de les Torres 4 con el mismo teléfono (fusionado)."),
    _pend_motivo="Venta de alimentación desconocida", _pend_info="Preguntar si vende pienso")

rec(nombre_comercial="Centre Veterinari Rubí Nord", tipo_establecimiento="veterinario", tipo_empresa="independiente",
    direccion="Carretera de Sabadell, 69", latitud=41.499266, longitud=2.032667, telefono="+34936687497",
    web="https://www.centroveterinariorubinord.es/",
    fuente_principal="Web oficial Centre Veterinari Rubí Nord",
    url_fuente="https://www.centroveterinariorubinord.es/contacto",
    fuentes_secundarias="; ".join([
        "https://www.qdq.com/centro-veterinario-rubi-nord-1231541",
        "https://peludosyfelices.com/veterinari-el-regne-animal-rubi/",
        ovt("Centre Veterinari Rubí Nord", "Meta, confianza 0,23"),
        ovt("El Reino Animal Veterinaria", "Meta; misma dirección, móvil 646 177 493")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Móvil 646 177 493. Horario oficial: L-J 9:30-13 y 17-20; V 9:30-13 y 16-19. En la "
                            "misma dirección figuraba antes 'Veterinari El Regne Animal' (fusionado como "
                            "duplicado). Sin evidencia de venta de alimentación."),
    _pend_motivo="Venta de alimentación desconocida", _pend_info="Preguntar si vende pienso")

rec(nombre_comercial="VetClínic 3 Calderón", tipo_establecimiento="veterinario", tipo_empresa="independiente",
    direccion="Carrer de Calderón de la Barca, 5", latitud=41.491507, longitud=2.036295, telefono="+34936978900",
    web="https://vetclinic.es/",
    fuente_principal="Web oficial VetClínic", url_fuente="https://vetclinic.es/vetclinic-3/",
    fuentes_secundarias="; ".join([
        "https://www.rubicomerc.cat/comercos.php?lang=ca&id=503",
        "https://www.paginasamarillas.es/f/rubi/vet-clinic-centros-veterinarios-s-l-_234271005_000000001.html",
        ovt("VetClinic Centros Veterinarios", "Meta, confianza 0,96")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Desde 2023 VetClínic concentra toda su actividad en este centro (web oficial). "
                            "Urgencias 608 99 77 10. Sus planes de salud mencionan descuentos en 'productos de "
                            "nuestra tienda' (indicio de tienda, sin confirmar alimentación)."),
    _pend_motivo="Venta de alimentación solo indiciaria", _pend_info="Confirmar si la tienda vende pienso")

rec(nombre_comercial="VetClínic 1 Les Torres", tipo_establecimiento="veterinario", tipo_empresa="independiente",
    direccion="Passeig de les Torres, 6, local 8", telefono="+34935884540",
    fuente_principal="Web oficial VetClínic", url_fuente="https://vetclinic.es/",
    fuentes_secundarias="https://veterinarioinfo7.com/c/rubi-08/; https://clinicaveterinaria.es/es-es/z/veterinarios/5264-rubi/",
    estado_registro="cerrado", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Cerrado: la web oficial indica que en 2023 VetClínic centralizó su actividad en "
                            "Calderón de la Barca 5. Se conserva para evitar reincorporarlo desde directorios "
                            "antiguos. Teléfono histórico de directorio."))

rec(nombre_comercial="VetClínic 2 Can Fatjó", tipo_establecimiento="veterinario", tipo_empresa="independiente",
    direccion="Carrer del Safir, 18", telefono="+34936974205",
    fuente_principal="Web oficial VetClínic", url_fuente="https://vetclinic.es/",
    fuentes_secundarias="https://clinicaveterinaria.es/es-es/i/23121-vet-clinic-centres-veterinaris-vetclinic-2-can-fatjo/; Overture 2026-09: en Safir 18 figura hoy 'Pavinou'",
    estado_registro="cerrado", nivel_verificacion="parcialmente_confirmado", metodo_verificacion="fuente oficial",
    observaciones_internas=("Cerrado: centralización de VetClínic en 2023 (web oficial). En la dirección figura "
                            "hoy otro negocio según Overture. Teléfono histórico de directorio."))

rec(nombre_comercial="Centre Veterinari Can Rosés", tipo_establecimiento="veterinario", tipo_empresa="independiente",
    direccion="Avinguda de les Olimpíades, 23-25", latitud=41.500036, longitud=2.03918, telefono="+34936979168",
    vende_alimentacion_mascotas="si",
    fuente_principal="Página de Facebook del centro",
    url_fuente="https://www.facebook.com/p/Centre-Veterinari-Can-Roses-100063630006388/",
    fuentes_secundarias="; ".join([
        "https://www.paginasamarillas.es/f/rubi/can-roses_181498619_000000001.html",
        "https://pininastyle.es/veterinario-en-rubi-centro-veterinario-can-roses/",
        "https://peludosyfelices.com/centro-veterinario-can-roses-rubi/",
        ovt("Centre Veterinari Can Roses", "Meta, confianza 0,93")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Varios directorios: clínica, venta de accesorios y alimentación, y peluquería. No "
                            "detallan especies. Overture le asigna la web barladrassana.com, que no corresponde "
                            "(descartada)."))

rec(nombre_comercial="Mª Assumpta Gaju Ricart (veterinaria)", tipo_establecimiento="veterinario",
    tipo_empresa="independiente", direccion="Carrer de Prat de la Riba, 17", telefono="+34935886941",
    fuente_principal="Citiservi", url_fuente="https://www.citiservi.es/barcelona/maria-assumpta-gaju-ricart-rubi__2848053_65.html",
    fuentes_secundarias="https://pininastyle.es/veterinario-en-rubi-maria-assumpta-gaju-ricart/; https://www.guiaveterinarios.com/pro/maria-assumpta-gaju-ricart-rubi-3976068/",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Solo directorios (uno cita 138 reseñas en Google, 4,8/5). No aparece en Overture "
                            "2026-09: sin coordenadas."),
    _pend_motivo="Solo directorios", _pend_info="Confirmar actividad, coordenadas y venta de alimentos")

rec(nombre_comercial="Consultorio Veterinario Can Oriol", tipo_establecimiento="veterinario", tipo_empresa=None,
    direccion="Carrer de Ciudad Real, 21", telefono="+34935885638",
    fuente_principal="Directorios veterinarios (resultado de búsqueda)",
    url_fuente="https://www.botiguesdecatalunya.cat/es/tiendas/veterinario/rubi",
    estado_registro="pendiente_revision", nivel_verificacion=NV,
    observaciones_internas=("Aparece en directorios (también citado como 'Canurion', no confirmado). No aparece "
                            "en Overture 2026-09."),
    _pend_motivo="Solo directorios", _pend_info="Confirmar existencia, nombre y coordenadas")

# ---------------------------------------------------------------- SUPERMERCADOS / HIPERMERCADOS
def super_(nombre, cadena, direccion, lat, lon, fuente, url, sec, metodo="fuente externa",
           estado="activo", nivel="parcialmente_confirmado", tel=None, web=None, obs="", tipo="supermercado",
           pend=None, verif=HOY):
    rec(nombre_comercial=nombre, tipo_establecimiento=tipo, tipo_empresa="cadena", nombre_cadena=cadena,
        direccion=direccion, latitud=lat, longitud=lon, telefono=tel, web=web,
        fuente_principal=fuente, url_fuente=url, fuentes_secundarias=sec, metodo_verificacion=metodo,
        estado_registro=estado, nivel_verificacion=nivel, fecha_ultima_verificacion=verif,
        observaciones_internas=(obs + " " + CHAIN_FOOD_NOTE).strip() if estado != "cerrado" else obs,
        _pend_motivo=pend[0] if pend else None, _pend_info=pend[1] if pend else None)


MERC_URL = "https://info.mercadona.es/es/supermercados?s=08191"
super_("Mercadona Avda. Castellbisbal", "Mercadona", "Avinguda de Castellbisbal, 58", 41.489965, 2.023084,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/mercadona-avda-castellbisbal/110477",
       "https://www.sucursales24.es/rubi/mercadona/avda-castellbisbal-58-1/; " + ovt("Mercadona", "Meta, confianza 1,0"),
       tel="+34933957705", web="https://www.mercadona.es/",
       obs="Teléfono de directorio. No figura en el extracto de AllThePlaces (localizador oficial) de Overture; "
           "sí en Meta y directorios. Páginas Amarillas cita también 'C/ Francesc Rovira 13' (no confirmada; "
           "posible dirección antigua).")
super_("Mercadona C/ de les Basses", "Mercadona", "Carrer de les Basses, 23", 41.495, 2.036094,
       "Localizador oficial Mercadona (vía AllThePlaces/Overture)",
       "https://info.mercadona.es/es/supermercados?s=08191&k=283185312920",
       "https://www.paginasamarillas.es/f/rubi/mercadona_012707303_000000572.html; " + ovt("Mercadona", "AllThePlaces"),
       metodo="fuente oficial", tel="+34935888251", web=MERC_URL,
       obs="Teléfono de Páginas Amarillas.")
super_("Mercadona C/ Marconi", "Mercadona", "Carrer de Marconi, 52", 41.482372, 2.040877,
       "Localizador oficial Mercadona (vía AllThePlaces/Overture)",
       "https://info.mercadona.es/es/supermercados?s=08191&k=283185316402",
       "https://www.sucursales24.es/rubi/mercadona/c-marconi-52-2/; " + ovt("Mercadona", "AllThePlaces"),
       metodo="fuente oficial", tel="+34934032135", web=MERC_URL, obs="Teléfono de directorio.")

super_("Lidl Avda. Electricitat", "Lidl", "Avinguda de l'Electricitat, 23", 41.47481, 2.032493,
       "Web oficial Lidl (ficha de tienda)", "https://www.lidl.es/s/es-ES/tiendas/rubi/avda-electricitat-23/",
       "https://www.revistaaral.com/texto-diario/mostrar/3906161/lidl-invierte-53-millones-segunda-tienda-rubi-crea-30-nuevos-empleos; " + ovt("Lidl", "Microsoft, confianza 0,98"),
       metodo="fuente oficial", web="https://www.lidl.es/s/es-ES/tiendas/rubi/avda-electricitat-23/",
       obs="Segunda tienda de Lidl en Rubí (2022). Algunos directorios usan el n.º 21. El 900 958 311 es "
           "atención al cliente de la cadena (no se registra como teléfono del establecimiento).")
super_("Lidl C/ Penedès", "Lidl", "Carrer del Penedès, 7", 41.484972, 2.041135,
       "Web oficial Lidl (tiendas en Rubí)", "https://www.lidl.es/s/es-ES/tiendas/rubi/",
       "https://www.supermercadocercademi.com/supermercados/lidl-rubi-carrer-del-penedes-7/; " + ovt("Lidl", "Meta, confianza 1,0"),
       metodo="fuente oficial", web="https://www.lidl.es/s/es-ES/tiendas/rubi/",
       obs="El listado oficial de Rubí muestra dos tiendas Lidl.")

super_("ALDI Rubí", "ALDI", "Avinguda de l'Estatut, 233", None, None,
       "Web oficial ALDI (ficha de tienda)",
       "https://www.aldi.es/encuentra-tu-supermercado/rubi/avinguda-de-l'estatut-233/3604738.html",
       "https://www.paginasamarillas.es/f/rubi/aldi_201200649_000000064.html; " + ovt("ALDI", "Microsoft; 41.495866, 2.038454") + "; " + ovt("ALDI Nord", "Meta; n.º 231; 41.502626, 2.037382"),
       metodo="fuente oficial",
       web="https://www.aldi.es/encuentra-tu-supermercado/rubi/avinguda-de-l'estatut-233/3604738.html",
       obs="Una sola tienda ALDI en Rubí. Overture tiene dos puntos separados unos 750 m (n.º 233 y n.º 231): "
           "coordenadas a NULL hasta comprobar en el mapa (la numeración vecina apunta a ~41.5026, 2.0374).",
       pend=("Coordenadas contradictorias", "Fijar coordenadas exactas"))

CONSUM_PRENSA = "https://www.consum.es/en/prensa/consum-opens-a-new-supermarket-in-rubi-and-exceeds-270-establishments-in-catalonia/"
super_("Consum Rubí Calderón de la Barca", "Consum", "Carrer de Calderón de la Barca, 6", 41.491409, 2.036537,
       "Web oficial Consum (ficha de tienda vía AllThePlaces/Overture)",
       "https://www.consum.es/supermercados/rubi-calderon-de-la-barca/",
       CONSUM_PRENSA + "; " + ovt("Consum", "AllThePlaces, confianza 0,95"),
       metodo="fuente oficial", tel="+34932768411",
       web="https://www.consum.es/supermercados/rubi-calderon-de-la-barca/",
       obs="Primer supermercado propio de Consum en Rubí (noviembre 2024), unos 1.500 m².")
for nombre, dire, lat, lon, url, sec, tel, obs, pend in [
    ("Charter Rubí Pl. de la Constitució", "Rambleta Joan Miró, s/n (Pl. de la Constitució)", 41.485396, 2.030502,
     "https://www.consum.es/supermercados/rubi-pl-de-la-constitucion/", ovt("Charter", "AllThePlaces"), None, "", None),
    ("Charter Rubí Virgen de Fátima", "Carrer de la Mare de Déu de Fàtima, 22", 41.486366, 2.036848,
     "https://www.consum.es/supermercados/rubi-virgen-de-fatima/",
     ovt("Supermercats Charter", "Microsoft") + "; " + ovt("Charter, la franquicia de Consum", "Meta; dirección errónea Av. Verge de Montserrat 22, misma URL"),
     "+34676302176", "Teléfono de Microsoft/Meta.", None),
    ("Charter Rubí Can Fatjó", "Avinguda de Can Fatjó, 1", 41.490044, 2.024755,
     "https://www.consum.es/en/supermarkets/rubi-avd-can-fatjo/", ovt("Supermercats Charter", "Microsoft, confianza 1,0"),
     "+34932661121", "Teléfono de Microsoft (sin contrastar).", None),
    ("Charter Rubí Carretera de Sabadell", "Carretera de Sabadell, 34", 41.498245, 2.032481,
     "https://www.consum.es/en/supermarkets/rubi-carretera-de-sabadell/", ovt("Supermercats Charter", "Microsoft"),
     "+34937443731", "Teléfono de Microsoft (sin contrastar).", None),
    ("Charter Rubí Mare de Déu de Lourdes", "Carrer de la Mare de Déu de Lourdes, 73", 41.49694, 2.040968,
     "https://www.consum.es/supermercados/rubi-mare-de-deu-de-lourdes/",
     ovt("Supermercats Charter", "Microsoft") + "; " + ovt("Consum", "Meta; n.º 75") + "; " + ovt("Charter", "Meta; dirección errónea Mare de Déu de Fàtima 61, misma URL"),
     "+34602469042", "Tres registros de Overture fusionados (n.º 73, n.º 75 y una dirección errónea con la misma URL oficial).", None),
    ("Charter Rubí Av. Castellnou", "Avinguda de Castellnou, 147, local 3", 41.536831, 1.991143,
     "https://www.consum.es/en/supermarkets/rubi-av-castellnou/",
     "https://www.paginasamarillas.es/f/rubi/charter-franquicias-de-consum_220591333_000000001.html; " + ovt("Consum", "Meta, confianza 1,0"),
     "+34935871826", "Urbanización Castellnou (dentro del término de Rubí según el polígono municipal).", None),
]:
    super_(nombre, "Consum (franquicia Charter)", dire, lat, lon, "Web oficial Consum (ficha de tienda)", url,
           CONSUM_PRENSA + "; " + sec, metodo="fuente oficial", tel=tel, web=url, obs=obs, pend=pend)

CONDIS = "Condis"
super_("Condis Can Fatjó", CONDIS, "Carrer del Diamant, 68", 41.488002, 2.024131,
       "Web oficial Condis (ficha de tienda)", "https://www.condis.es/ficha/2713",
       "https://www.tiendeo.com/Tiendas/rubi/condis-c-del-diamant/159815; " + ovt("Condis Can Fatjó", "Meta, confianza 1,0"),
       metodo="fuente oficial", tel="+34936972075", web="https://www.condis.es/ficha/2713",
       obs="La ficha oficial da el correo genérico info@condis.es (no se registra como email del establecimiento).")
super_("Condis Duran i Bas", CONDIS, "Carrer de Duran i Bas, 12", 41.487416, 2.036596,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/condis-c-duran-i-bas/1458",
       "https://sucursales.cylex.es/rubi/supermercados/Condis/; " + ovt("Condis el Teu Super", "Microsoft, confianza 0,97"),
       tel="+34935887548", web="https://www.condis.es/")
super_("Condis Mercat de Rubí", CONDIS, "Carrer de Cal Gerrer, 1 (Mercat Municipal)", 41.48934, 2.03479,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/condis-c-cal-gerrer-mercat-rubi--rubi/159817",
       ovt("CondisLife el Teu Super", "Meta; C/ Cal Gerrer 7") + "; Overture: Amazon Locker 'Supermercado Condis Mercat, Carrer de Cal Gerrer 1'",
       tel="+34936991515", web="https://www.condis.es/",
       obs="Overture lo recoge como 'CondisLife' en Cal Gerrer 7 (fusionado). Teléfono de Meta.")
super_("Condis Avinguda Barcelona", CONDIS, "Avinguda de Barcelona, 44-46", 41.488118, 2.030971,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/condis-avda-barcelona/1590",
       "https://sucursales.cylex.es/rubi/supermercados/Condis/; " + ovt("Condis", "Microsoft, confianza 0,97"),
       tel="+34935860043", web="https://www.condis.es/")
super_("Condis Express Pau Claris", CONDIS, "Passeig de Pau Claris, 63", 41.496493, 2.031192,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/condis-pi-pau-claris/249718",
       ovt("Condis Express", "Microsoft, confianza 0,97"), tel="+34931778426", web="https://www.condis.es/")
super_("Condis Express Carretera de Sabadell", CONDIS, "Carretera de Sabadell, 89", 41.500375, 2.032968,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/condis-c-sabadell/159829",
       ovt("Condis Express", "Meta, confianza 1,0"), web="https://www.condis.es/",
       obs="Teléfono contradictorio (Tiendeo 935 88 01 18; Meta 937 63 87 93): se deja NULL.",
       pend=("Teléfono contradictorio", "Confirmar teléfono"))
super_("Condis Can Cabanyes (Torre Iris)", CONDIS, "Avinguda de Can Cabanyes, s/n (Torre Iris)", 41.485475, 2.031908,
       "Localizador oficial Condis (vía AllThePlaces/Overture)", "https://www.condis.es/ficha/3837",
       ovt("Condis", "AllThePlaces"), metodo="fuente externa", estado="pendiente_revision",
       web="https://www.condis.es/ficha/3837",
       obs="Solo una fuente (extracto del localizador oficial en Overture); no contrastado en web.",
       pend=("Una sola fuente", "Confirmar actividad"))
super_("Condis C/ Sabadell 41", CONDIS, "Carrer de Sabadell, 41", 41.496184, 2.034452,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/condis",
       ovt("Condis", "Meta; URL de Consum") + "; " + ovt("Top Rubi", "Meta; mismo teléfono 936 97 18 02"),
       estado="pendiente_revision", nivel=NV, tel="+34936971802",
       obs="Registro confuso: Overture mezcla Condis, una URL de Charter (Ctra. de Sabadell) y otro nombre "
           "('Top Rubi') con el mismo teléfono. Puede haber cambiado de enseña.",
       pend=("Enseña y actividad contradictorias", "Confirmar enseña actual"))
super_("Condis Joan Maragall", CONDIS, "Carrer de Joan Maragall, 15", None, None,
       "Páginas Amarillas", "https://www.paginasamarillas.es/f/rubi/condis_012008702_000000086.html",
       "https://www.cylex.es/rubi/condis-12169323.html", estado="pendiente_revision", nivel=NV,
       obs="Solo directorios; no aparece en Overture 2026-09.", pend=("Solo directorios", "Confirmar existencia y coordenadas"))

DIA_URL = "https://www.dia.es/tiendas/buscador-tiendas/barcelona/rubi"
for nombre, dire, lat, lon, sec in [
    ("Dia Mestre Magí Ramentol", "Carrer del Mestre Magí Ramentol, 17", 41.490014, 2.034775, ovt("Supermercados Dia", "Meta, confianza 1,0")),
    ("Dia Avinguda de Barcelona", "Avinguda de Barcelona, s/n", 41.499842, 2.040488,
     ovt("DIA España", "Meta") + "; " + ovt("Dia", "Microsoft; 41.499675, 2.040598") + "; https://www.waze.com/live-map/directions/es/ct/rubi/supermercats-dia?to=place.ChIJLU9xu2ORpBIRn1WmfxsskOg"),
    ("Dia Pitàgores", "Carrer de Pitàgores, 5", 41.485301, 2.035852, ovt("Supermercados Dia", "Meta, confianza 0,82")),
    ("Dia Pau Claris", "Passeig de Pau Claris, 56", 41.495435, 2.031059, ovt("Supermercados Dia", "Meta, confianza 0,47")),
]:
    super_(nombre, "Dia", dire, lat, lon, "Buscador oficial de tiendas Dia", DIA_URL, sec,
           metodo="fuente oficial", web=DIA_URL,
           obs="El 912 170 453 de las fichas es un número corporativo (no se registra).")

super_("Caprabo Pau Claris", "Caprabo", "Passeig de Pau Claris, 38-40", 41.494899, 2.031088,
       "Food Retail (noticia de apertura)", "https://www.foodretail.es/retailers/Caprabo-supermercados-Rubi-Badalona-Girona_0_1502549743.html",
       "https://ahorra-ya.es/tiendas/rubi/caprabo/passeig-pau-claris-38-40; " + ovt("Caprabo", "Microsoft"),
       web="https://www.caprabo.com/",
       obs="Unos 300 m². El 932 616 060 es el número corporativo de Caprabo (no se registra).")
super_("Caprabo Plaça Esbart Dansaire (Aliprox)", "Caprabo (franquicia Aliprox)", "Plaça de l'Esbart Dansaire, 5",
       41.497372, 2.035649, "Ràdio Rubí (noticia de apertura)",
       "https://www.radiorubi.cat/news/caprabo-obre-un-nou-supermercat-a-la-serreta/",
       "https://www.foodretail.es/retailers/caprabo-crece-en-la-provincia-de-barcelona-con-un-nuevo-supermercado-aliprox-en-rubi.html; " + ovt("Caprabo", "Foursquare"),
       web="https://www.caprabo.com/",
       obs="Franquicia Aliprox de unos 200 m², junto al institut La Serreta; L-D 8-23 según la noticia.")

super_("bonÀrea Mare de Déu de Lourdes", "bonÀrea", "Carrer de la Mare de Déu de Lourdes, 15", 41.494703, 2.037424,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/bonarea-cl-ntra-senyora-de-lourdes/40073",
       "https://www.paginasamarillas.es/f/rubi/bonarea_179403662_000000204.html; https://www.cylex.es/rubi/bon%C3%A0rea-11328494.html; " + ovt("Bonàrea", "Foursquare"),
       tel="+34935889736", web="https://www.bonarea.com/")
super_("bonÀrea Llobateras", "bonÀrea", "Carrer de Llobateras, 92", 41.488917, 2.035023,
       "Waze", "https://www.waze.com/live-map/directions/es/ct/rubi/bonarea?to=place.ChIJlaZ07F6RpBIRz_4egTdj87g",
       ovt("bonÀrea", "Meta, confianza 0,89"), tel="+34936815719", web="https://www.bonarea.com/",
       obs="Dos fuentes (Waze y Meta). Hay otras direcciones antiguas de bonÀrea en directorios (C/ Lluís Ribas 27, "
           "C/ Narcís Menard 9-11) no confirmadas.")

super_("Esclat Rubí", "BonpreuEsclat (Grup Bon Preu)", "Carretera de Molins de Rei, 75", 41.47699, 2.03138,
       "Web oficial BonpreuEsclat (ficha de tienda)", "https://www.bonpreuesclat.cat/es/esclat-rub%C3%AD",
       "https://www.diariderubi.com/bon-preu-ha-inaugurat-un-esclat-a-rubi/; https://www.rubitv.cat/societat/sinaugura-el-primer-esclat-a-rubi_943341_102.html; " + ovt("Esclat", "Microsoft, confianza 0,98"),
       metodo="fuente oficial", tel="+34937617528", web="https://www.bonpreuesclat.cat/es/esclat-rub%C3%AD",
       tipo="hipermercado",
       obs="Abierto en junio de 2020 con 3.058 m² de sala de ventas (>2.500 m²: se clasifica como hipermercado). "
           "L-S 9-21.")
super_("Alcampo Rubí", "Alcampo (Auchan)", "Carrer de Miguel de Cervantes, 79", 41.48832, 2.03362,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/alcampo-calle-cervantes/264702",
       "https://kosmosmarket.com/tienda/alcampo-811; " + ovt("Alcampo", "Meta, confianza 1,0") + "; " + ovt("Mi Alcampo", "Foursquare (duplicado)"),
       web="https://www.alcampo.es/",
       obs="Formato supermercado urbano (no hipermercado; superficie no confirmada). El 913 687 857 es el "
           "número corporativo.")
super_("Sorli Rubí", "Sorli", "Carrer de Monturiol, 28", 41.48599, 2.033965,
       "Tiendeo", "https://www.tiendeo.com/Tiendas/rubi/sorli-discau-carrer-monturiol/247350",
       "https://www.facebook.com/sorlirubi/; https://www.ubereats.com/es/store/sorli-monturiol-rubi/PtP1GWRvX4iPJylt2vG6aA; " + ovt("Sorli", "Meta"),
       tel="+34935610337", web="http://www.sorlidiscau.es",
       obs="Horarios contradictorios entre directorios.")
super_("Carrefour Express Sant Joan", "Carrefour Express", "Carrer de Sant Joan, 12-14", 41.490586, 2.030316,
       "Web oficial Carrefour (ficha de tienda)",
       "https://www.carrefour.es/tiendas-carrefour/supermercados/carrefour-express/calle-sant-joan-rubi-s.aspx",
       ovt("Carrefour Express Calle Sant Joan - Rubí", "AllThePlaces"), metodo="fuente oficial",
       web="https://www.carrefour.es/tiendas-carrefour/supermercados/carrefour-express/calle-sant-joan-rubi-s.aspx",
       obs="L-S 9-21:30 según la ficha oficial.")
super_("Carrefour Express Joan Maragall", "Carrefour Express", "Carrer de Joan Maragall, 29", 41.48806, 2.036857,
       "Web oficial Carrefour (ficha de tienda)",
       "https://www.carrefour.es/tiendas-carrefour/supermercados/carrefour-express/joan-maragall-rubi-s.aspx",
       "https://radarsuper.com/c/barcelona/rubi/carrefour-joan-maragall-rubi-s; " + ovt("Carrefour Express Joan Maragall - Rubí", "AllThePlaces"),
       metodo="fuente oficial", estado="cerrado",
       obs="La ficha oficial de Carrefour la marca como 'Cerrada'. Directorios la siguen mostrando abierta.")
super_("SUMA Supermercat Bartrina", "SUMA", "Carrer de Bartrina, 22", 41.488091, 2.037936,
       "Waze", "https://www.waze.com/es/live-map/directions/es/ct/rubi/suma-supermercado?to=place.ChIJNz2YHh6RpBIRU0dpnlh-cAs",
       "https://www.tiendeo.com/Tiendas/rubi/suma-supermercados-calle-bartrina/251089; " + ovt("SUMA Supermercat", "Microsoft"),
       tel="+34655016786", web="https://www.sumasupermercados.es/ca",
       obs="L-D 9-23 según Waze. Tiendeo cita otro SUMA en C/ Monturiol 48 (no confirmado; no incorporado).")
super_("Covirán Rubí", "Covirán", "Carrer de la Mare de Déu de Fàtima, 18", 41.486893, 2.035536,
       "Localizador oficial Covirán (vía AllThePlaces/Overture)", "https://www.coviran.es/",
       ovt("Covirán", "AllThePlaces"), estado="pendiente_revision",
       obs="Solo una fuente; la búsqueda web no la confirma.", pend=("Una sola fuente", "Confirmar existencia"))

# Supermercados independientes de barrio (solo Overture/Meta, sin contrastar)
for nombre, dire, lat, lon, tel, oid, obs in [
    ("Ali SuperMercat", "Plaça del Progrès, 9", 41.486158, 2.036257, "+34632115081", "Ali SuperMercat (Meta, confianza 0,69)", ""),
    ("Broas Supermercado", "Passeig de les Torres, 62", 41.48105, 2.033301, "+34651829224", "Broas supermercado (Meta, confianza 0,51)",
     "Comparte teléfono con 'El Rebost del Barri'."),
    ("El Rebost del Barri", "Carretera de Terrassa, 78", 41.500134, 2.032477, "+34651829224", "El Rebost del Barri (Meta, confianza 0,5)",
     "Comparte teléfono con 'Broas Supermercado'."),
    ("Supermercat (C/ Lepant)", "Carrer de Lepant, 61", 41.500934, 2.035522, None, "Supermercat (Meta, confianza 0,71)",
     "Sin nombre comercial en la fuente."),
    ("El Barri La Plana", "Carrer de la Plana de Can Bertran, 29", 41.493362, 2.035668, "+34935881917", "El Barri La Plana (Meta, confianza 0,94)",
     "Tienda de conveniencia."),
]:
    rec(nombre_comercial=nombre, tipo_establecimiento="supermercado", tipo_empresa="independiente",
        direccion=dire, latitud=lat, longitud=lon, telefono=tel,
        fuente_principal=OVT, url_fuente="https://docs.overturemaps.org/guides/places/",
        fuentes_secundarias=f"Overture: {oid}", fecha_ultima_verificacion=None, metodo_verificacion=None,
        estado_registro="pendiente_revision", nivel_verificacion=NV,
        observaciones_internas=("Supermercado de barrio independiente; solo una fuente cartográfica, sin contraste "
                                "web. " + obs).strip(),
        _pend_motivo="Una sola fuente", _pend_info="Confirmar existencia, actividad y surtido de mascotas")

# ---------------------------------------------------------------- GARDEN
rec(nombre_comercial="Centre de Jardineria La Noguera", tipo_establecimiento="garden", tipo_empresa="independiente",
    direccion="Carretera de Molins de Rei a Rubí, km 9,8", latitud=41.478951, longitud=2.030953,
    telefono="+34935872001", web="http://www.lanoguera.es/", vende_alimentacion_mascotas="si",
    fuente_principal="El Comerç de Rubí (rubicomerc.cat)",
    url_fuente="https://www.rubicomerc.cat/comercos.php?lang=ca&id=585",
    fuentes_secundarias="; ".join([
        "https://www.walk2view.com/centre_de_jardineria_la_noguera_rubi_barcelona_espana-142300-en.html",
        "https://www.cylex.es/rubi/centre-de-jardineria-la-noguera-10968716.html",
        "https://twitter.com/gardenlanoguera", ovt("Centre de Jardineria La Noguera", "Meta, confianza 1,0")]),
    estado_registro="activo", nivel_verificacion="parcialmente_confirmado",
    observaciones_internas=("Garden de unos 8.000 m², socio fundador de Jardinarium (asociación, no cadena). "
                            "Sección de mascotas: 'complementos y alimentos para tu mascota' (walk2view) y "
                            "accesorios (rubicomerc). Km 9,5 o 9,8 según la fuente. Overture tiene otro punto "
                            "'La Noguera' unos 600 m al sur (descartado). Especies no detalladas."))

# ---------------------------------------------------------------- TIENDAS NATURALES
for nombre, dire, lat, lon, tel, web, fuente, url, sec, obs, nivel, estado in [
    ("Herbolari de Rubí", "Carrer de Magallanes, 6", 41.491828, 2.033528, "+34936974865", "https://herbolariderubi.es/",
     "El Comerç de Rubí (rubicomerc.cat)", "https://www.rubicomerc.cat/comercos.php?id=443&lang=ca",
     "https://www.qdq.com/herbolari-de-rubi-454935; " + ovt("Herbolario de Rubí", "Meta, confianza 0,95"),
     "Herbolario, dietética y productos sin gluten.", "parcialmente_confirmado", "activo"),
    ("Herbo-Dietètica L'Estel", "Carrer de Prat de la Riba, 25", 41.494533, 2.03164, "+34935883846", None,
     "Botigues de Catalunya", "https://www.botiguesdecatalunya.cat/es/tiendas/herbolario/rubi",
     "rubicomerc.cat (búsqueda 'herboristería'); " + ovt("Herbo-Dietetica L'Estel", "Meta, confianza 0,46"),
     "Directorio cita a 'Coral López García' en la misma dirección.", "parcialmente_confirmado", "activo"),
    ("Herbodietètica Rosa Gustems", "Carrer de Monturiol, 24", 41.486301, 2.034002, "+34936975749", None,
     "Botigues de Catalunya", "https://www.botiguesdecatalunya.cat/es/tiendas/herbolario/rubi",
     ovt("Herbodietetica Rosa Gustems", "Microsoft, confianza 0,85"), "", "parcialmente_confirmado", "activo"),
    ("Sentits de Vida", "Carrer del Doctor Maximí Fornés, 44", 41.491399, 2.030954, "+34935873188", None,
     "El Comerç de Rubí (rubicomerc.cat)", "https://www.rubicomerc.cat/resultats.php?lang=ca&search_type=products&search_term=herborister%C3%ADa",
     ovt("Sentits de Vida", "Meta, confianza 0,81"), "Overture la clasifica como medicina naturista.",
     "parcialmente_confirmado", "activo"),
    ("Herboristeria Majó", "Carrer de Sant Pere, 66", None, None, None, None,
     "QDQ", "https://www.qdq.com/herboristeria-majo-771500",
     "https://www.botiguesdecatalunya.cat/es/tiendas/herbolario/rubi",
     "Con otra parada en el Mercat Municipal (puestos 41-42). No localizada en Overture: sin coordenadas.",
     NV, "pendiente_revision"),
]:
    rec(nombre_comercial=nombre, tipo_establecimiento="tienda_natural", tipo_empresa="independiente",
        direccion=dire, latitud=lat, longitud=lon, telefono=tel, web=web, fuente_principal=fuente, url_fuente=url,
        fuentes_secundarias=sec, estado_registro=estado, nivel_verificacion=nivel,
        observaciones_internas=(obs + " Sin evidencia de venta de alimentación para mascotas.").strip(),
        _pend_motivo="Venta de alimentación para mascotas desconocida" if estado == "activo" else "Solo directorios",
        _pend_info="Preguntar si venden alimento o snacks naturales para mascotas")


# ---------------------------------------------------------------- DESCARTADOS (hoja auxiliar)
DESCARTADOS = [
    ("Tiendanimal Sant Cugat", "Ctra. de Sant Cugat a Rubí (junto a Caprabo), 08174", "Fuera de Rubí",
     "Localizador oficial la sitúa en Sant Cugat (CP 08174, URL /sant-cugat-del-valles/); el punto de Overture queda ~70 m fuera del polígono de Rubí.",
     "https://tiendas.tiendanimal.es/cataluna/barcelona/sant-cugat-del-valles/tiendanimal-sant-cugat-barcelona-12"),
    ("Caprabo Pol. Can Sant Joan, parcel·la 10", "Pol. Ind. Can Sant Joan, Sant Cugat", "Fuera de Rubí",
     "Páginas Amarillas y Tiendeo lo sitúan en Sant Cugat del Vallès; punto de Overture fuera del polígono.",
     "https://www.tiendeo.com/Tiendas/sant-cugat-del-valles/caprabo-pol-ind-can-st-joan-parcel%C2%B7la-costat-ctra-rubi/198354"),
    ("Leroy Merlin Sant Cugat", "Carretera de Rubí, 7, 08174 Sant Cugat", "Fuera de Rubí",
     "Web oficial y nota del INCASÒL: Parc d'activitats econòmiques Can Sant Joan, término de Sant Cugat.",
     "https://www.leroymerlin.es/tiendas/sant-cugat.html"),
    ("Tiendanimal Terrassa", "Pol. Santa Margarida II, 08223 Terrassa", "Fuera de Rubí", "Terrassa.", None),
    ("Bricomart", "Terrassa / Sant Quirze", "Fuera de Rubí", "Punto de Overture fuera del polígono.", None),
    ("Kiwoko", "-", "Sin tienda en Rubí", "Localizador oficial: no hay tienda en Rubí ni en Sant Cugat.", "https://tiendas.kiwoko.com/catalunya"),
    ("Vivers Camp-Rubí", "Ctra. Santa Creu de Calafell km 3,8, Gavà", "Fuera de Rubí", "Confusión de nombre (está en Gavà).", "https://camprubijardins.com/"),
    ("Visa Garden", "Carrer Anoia, 2, Rubí", "Fuera de alcance", "Proveedor profesional (B2B), sin venta al público.", "https://visa-garden.com/en/"),
    ("Rubinum", "Av. de la Llana, 123, Rubí", "Fuera de alcance", "Laboratorio o industria farmacéutica, sin venta al público.", None),
    ("Centro Comercial Sant Cugat del Vallès", "-", "Registro erróneo", "Overture lo sitúa en Rubí por error (el centro está en Sant Cugat).", None),
    ("Dia Market C/ Saragossa", "Carrer de Saragossa, Rubí", "No confirmado", "No figura en el buscador oficial de Dia (4 tiendas en Rubí). Posible registro antiguo.", None),
    ("Minimercat (Ctra. Molins de Rei 69)", "Rubí", "No confirmado", "Registro con web de BonpreuEsclat junto al Esclat; probable tienda de gasolinera o duplicado.", None),
    ("SUMA C/ Monturiol 48", "Rubí", "No confirmado", "Solo en Tiendeo.", None),
    ("bonÀrea C/ Lluís Ribas 27 / C/ Narcís Menard 9-11", "Rubí", "No confirmado", "Direcciones de directorios antiguos.", None),
    ("Mercadona C/ Francesc Rovira 13", "Rubí", "No confirmado", "Solo en Páginas Amarillas; posible dirección antigua.", None),
    ("Residència canina Can Xavi", "Urb. Can Carreras, 47, Rubí", "Fuera de alcance V1", "Residencia canina (servicio); sin evidencia de venta.", "http://www.canxavi.cat/"),
    ("Rott de Rubí", "Rubí", "Fuera de alcance V1", "Criadero; sin establecimiento de venta.", None),
    ("Naturhouse Rubí", "Carrer de la Creu, 5, Rubí", "Fuera de alcance", "Centro de dietética humana.", None),
    ("Herbolari Gaia, Santiveri (Pg. Francesc Macià 28), Yerbabuena (C/ Bartrina 1), Herbosalut (C/ General Castaños 31), L'Artemisa (C/ Terrassa 6), Natur Rubí (C/ Creu 5), El Rebost de la Natura",
     "Rubí", "No incorporados (evidencia insuficiente)",
     "Solo una fuente de directorio sin fecha, o no localizados en Overture. Candidatos para una revisión posterior.",
     "https://www.botiguesdecatalunya.cat/es/tiendas/herbolario/rubi"),
    ("Dani Supermercat", "Sin dirección", "No incorporado", "Registro de Overture sin dirección.", None),
]

DUPLICADOS = [
    ("El Reino Animal Veterinaria (Ctra. Sabadell 69)", "Centre Veterinari Rubí Nord"),
    ("Las Torres (veterinario, Pg. Torres 4)", "Centre Veterinari Les Torres"),
    ("SonsoVet Homeopatia (C/ Gimnàs 17)", "Dispensari Veterinari del Vallès - Rubí 2"),
    ("Tienda especializada en piensos... (Pg. Torres 49)", "Boutique Animal"),
    ("DoggyRade", "Sweet Doggy"),
    ("ALDI Nord (Estatut 231)", "ALDI Rubí"),
    ("Mi Alcampo", "Alcampo Rubí"),
    ("Consum (Lourdes 75)", "Charter Rubí Mare de Déu de Lourdes"),
    ("Charter (dirección errónea Mare de Déu de Fàtima 61)", "Charter Rubí Mare de Déu de Lourdes"),
    ("Charter (dirección errónea Av. Verge de Montserrat 22)", "Charter Rubí Virgen de Fátima"),
    ("Dia (Avinguda de Barcelona s/n, 2.º registro)", "Dia Avinguda de Barcelona"),
    ("CondisLife el Teu Super (Cal Gerrer 7)", "Condis Mercat de Rubí"),
    ("Top Rubi (Sabadell 41)", "Condis C/ Sabadell 41"),
    ("La Noguera (2.º punto de Overture)", "Centre de Jardineria La Noguera"),
]


def validar():
    for r in R:
        for k, allowed in PERMITIDOS.items():
            assert r[k] in allowed, (r["nombre_comercial"], k, r[k])
        if r["tipo_empresa"] != "cadena":
            assert r["nombre_cadena"] is None, r["nombre_comercial"]
        if r["vende_alimentacion_mascotas"] != "si":
            assert "si" not in (r["perros"], r["gatos"], r["otros_animales"]), r["nombre_comercial"]
        if r["nivel_verificacion"] == "confirmado":
            assert r["vende_alimentacion_mascotas"] == "si" and r["estado_registro"] == "activo"
        assert (r["latitud"] is None) == (r["longitud"] is None)
        if r["latitud"] is not None:
            assert 41.46 < r["latitud"] < 41.545 and 1.97 < r["longitud"] < 2.06, r["nombre_comercial"]
    claves = Counter((r["nombre_comercial"].lower(), r["direccion"].lower()) for r in R)
    assert max(claves.values()) == 1
    coords = Counter((r["latitud"], r["longitud"]) for r in R if r["latitud"] is not None)
    assert max(coords.values()) == 1, [c for c, n in coords.items() if n > 1]


ORDEN_TIPO = ["tienda_mascotas", "agrotienda", "veterinario", "garden", "hipermercado", "supermercado",
              "tienda_natural", "otros"]


def asignar_ids(fichero_ids):
    """Los identificadores son estables: se leen de ids_rubi_v1.csv (clave interna = nombre|dirección del
    primer lote) y solo se asignan números nuevos, nunca reutilizados, a los registros que no estén."""
    R.sort(key=lambda r: (ORDEN_TIPO.index(r["tipo_establecimiento"]), r["estado_registro"] == "cerrado"))
    previos = {}
    if fichero_ids.exists():
        with open(fichero_ids, encoding="utf-8") as fh:
            previos = {row["clave"]: row["id_zimia"] for row in csv.DictReader(fh)}
    siguiente = max([int(v.rsplit("-", 1)[1]) for v in previos.values()] + [0]) + 1
    for r in R:
        clave = r.get("_clave") or f"{r['nombre_comercial']}|{r['direccion']}"
        r["_clave"] = clave
        if clave not in previos:
            previos[clave] = f"ZIM-RUB-{siguiente:04d}"
            siguiente += 1
        r["id_zimia"] = previos[clave]
    with open(fichero_ids, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["id_zimia", "clave"])
        for clave, i in sorted(previos.items(), key=lambda kv: kv[1]):
            w.writerow([i, clave])
    R.sort(key=lambda r: r["id_zimia"])


HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def hoja(ws, cabecera, filas, anchos=None):
    ws.append(cabecera)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "B2"
    for i, h in enumerate(cabecera, 1):
        ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(h, min(max(len(h) + 2, 12), 45))
    if filas:
        ws.auto_filter.ref = ws.dimensions


def pct(n, d):
    return f"{(100 * n / d):.1f}%" if d else "0%"


def main(outdir):
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    validar()
    asignar_ids(Path(__file__).resolve().parent / "ids_rubi_v1.csv")
    n = len(R)

    # ---- CSV (NULL -> vacío)
    with open(out / "Zimia_Establecimientos_Rubi_V1.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(CAMPOS)
        for r in R:
            w.writerow(["" if r[c] is None else r[c] for c in CAMPOS])

    wb = Workbook()
    ws = wb.active
    ws.title = "ESTABLECIMIENTOS"
    anchos = {"nombre_comercial": 38, "direccion": 40, "fuente_principal": 40, "url_fuente": 50,
              "fuentes_secundarias": 60, "observaciones_internas": 80, "web": 40}
    hoja(ws, CAMPOS, [[("NULL" if r[c] is None else r[c]) for c in CAMPOS] for r in R], anchos)

    # ---- Pendientes
    pend = [r for r in R if r["estado_registro"] != "cerrado"
            and (r["_pend_motivo"] or r["nivel_verificacion"] != "confirmado")]
    ws2 = wb.create_sheet("PENDIENTES DE VERIFICACIÓN")
    cab2 = ["id_zimia", "nombre_comercial", "tipo_establecimiento", "direccion", "estado_registro",
            "nivel_verificacion", "motivo_revision", "informacion_pendiente", "campos_sin_dato",
            "fuentes_consultadas", "prioridad"]
    filas2 = []
    for r in pend:
        faltan = [c for c in ("latitud", "telefono", "web", "email") if r[c] is None]
        if r["vende_alimentacion_mascotas"] != "si":
            faltan.append("vende_alimentacion_mascotas")
        elif NV in (r["perros"], r["gatos"]):
            faltan.append("perros/gatos")
        motivo = r["_pend_motivo"] or ("Venta de alimentación no verificada en la sucursal"
                                       if r["tipo_establecimiento"] in ("supermercado", "hipermercado")
                                       else "Información incompleta")
        info = r["_pend_info"] or "Confirmar surtido de alimentación para perros y gatos"
        prio = "alta" if r["tipo_establecimiento"] in ("tienda_mascotas", "agrotienda", "veterinario") else (
            "media" if r["tipo_establecimiento"] in ("garden", "hipermercado", "otros") else "baja")
        filas2.append([r["id_zimia"], r["nombre_comercial"], r["tipo_establecimiento"], r["direccion"],
                       r["estado_registro"], r["nivel_verificacion"], motivo, info, ", ".join(faltan),
                       "; ".join(x for x in (r["url_fuente"], r["fuentes_secundarias"]) if x), prio])
    hoja(ws2, cab2, filas2, {"nombre_comercial": 38, "direccion": 38, "motivo_revision": 45,
                             "informacion_pendiente": 50, "fuentes_consultadas": 80, "campos_sin_dato": 35})

    # ---- Resumen
    ws3 = wb.create_sheet("RESUMEN")
    act = [r for r in R if r["estado_registro"] != "cerrado"]
    na = len(act)
    filas3 = [["Fecha de la investigación", HOY, ""],
              ["Total de registros en la base", n, "Incluye cerrados conservados por trazabilidad"],
              ["Registros activos o pendientes (no cerrados)", na, ""],
              ["Registros cerrados confirmados", n - na, ""],
              ["", "", ""], ["POR CATEGORÍA (no cerrados)", "n.º", "%"]]
    for t in ORDEN_TIPO:
        k = sum(1 for r in act if r["tipo_establecimiento"] == t)
        filas3.append([t, k, pct(k, na)])
    filas3 += [["", "", ""], ["POR TIPO DE EMPRESA (no cerrados)", "n.º", "%"]]
    for t in ("independiente", "cadena", None):
        k = sum(1 for r in act if r["tipo_empresa"] == t)
        filas3.append([t or "NULL (no determinado)", k, pct(k, na)])
    filas3 += [["", "", ""], ["NIVEL DE VERIFICACIÓN (todos)", "n.º", "%"]]
    for t in ("confirmado", "parcialmente_confirmado", "no_verificado"):
        k = sum(1 for r in R if r["nivel_verificacion"] == t)
        filas3.append([t, k, pct(k, n)])
    filas3 += [["", "", ""], ["ESTADO DEL REGISTRO (todos)", "n.º", "%"]]
    for t in ("activo", "pendiente_revision", "cerrado"):
        k = sum(1 for r in R if r["estado_registro"] == t)
        filas3.append([t, k, pct(k, n)])
    filas3 += [["", "", ""], ["ALIMENTACIÓN (no cerrados)", "n.º", "%"]]
    for campo, lab in (("vende_alimentacion_mascotas", "Venta de alimentación para mascotas confirmada (si)"),
                       ("perros", "Alimentación para perros confirmada"),
                       ("gatos", "Alimentación para gatos confirmada"),
                       ("otros_animales", "Alimentación para otros animales confirmada")):
        k = sum(1 for r in act if r[campo] == "si")
        filas3.append([lab, k, pct(k, na)])
    k = sum(1 for r in act if r["vende_alimentacion_mascotas"] == NV)
    filas3.append(["Venta de alimentación no verificada", k, pct(k, na)])
    filas3 += [["", "", ""], ["COBERTURA DE CAMPOS (todos los registros)", "n.º con dato", "% cobertura"]]
    for c in ("direccion", "latitud", "telefono", "web", "email", "tipo_empresa", "fecha_ultima_verificacion"):
        k = sum(1 for r in R if r[c] is not None)
        filas3.append([c if c != "latitud" else "coordenadas (latitud/longitud)", k, pct(k, n)])
    filas3 += [["", "", ""], ["CONTROL DE CALIDAD", "n.º", ""],
               ["Registros en hoja de pendientes", len(filas2), pct(len(filas2), n)],
               ["Posibles duplicados eliminados (fusionados)", len(DUPLICADOS), "Ver hoja DESCARTADOS"],
               ["Candidatos descartados o no incorporados", len(DESCARTADOS), "Ver hoja DESCARTADOS"],
               ["Puntos de Overture analizados dentro del polígono de Rubí", 1854, "Release 2026-09-23.0"],
               ["Candidatos preseleccionados por categoría o palabra clave", 255, "Antes de la revisión manual"]]
    hoja(ws3, ["Indicador", "Valor", "Detalle"], filas3, {"Indicador": 60, "Valor": 16, "Detalle": 45})

    # ---- Metodología
    ws4 = wb.create_sheet("METODOLOGÍA")
    for linea in METODOLOGIA:
        ws4.append(linea)
    ws4.column_dimensions["A"].width = 32
    ws4.column_dimensions["B"].width = 140
    for row in ws4.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[0].value and not (len(row) > 1 and row[1].value):
            row[0].font = Font(bold=True)

    # ---- Descartados y duplicados (auxiliar)
    ws5 = wb.create_sheet("DESCARTADOS")
    hoja(ws5, ["candidato", "direccion", "motivo", "detalle", "fuente"],
         [list(d) for d in DESCARTADOS] + [["", "", "", "", ""], ["DUPLICADO FUSIONADO", "", "", "registro conservado", ""]]
         + [[a, "", "Duplicado", b, OVT] for a, b in DUPLICADOS],
         {"candidato": 50, "direccion": 40, "motivo": 28, "detalle": 80, "fuente": 60})

    wb.save(out / "Zimia_Establecimientos_Rubi_V1.xlsx")
    return {"total": n, "pendientes": len(filas2)}


METODOLOGIA = [
    ["1. FUENTES CONSULTADAS"],
    ["Cartografía base", "Overture Maps Foundation, release 2026-09-23.0, temas 'places' y 'divisions' (s3://overturemaps-us-west-2). "
     "Places agrega datos de Meta, Microsoft, Foursquare, AllThePlaces (extracción de localizadores oficiales de cadenas) y PinMeTo. "
     "Límite municipal: división 'Rubí' (locality, id 7f6eeace-47c8-41b9-bf66-82185f798ef8), polígono derivado de OpenStreetMap, "
     "área ≈ 32 km² (coincide con la superficie oficial de Rubí, 32,3 km²)."],
    ["Webs oficiales y localizadores", "Lidl, ALDI, Dia, Consum/Charter, Condis, Carrefour, BonpreuEsclat, Mercadona (vía AllThePlaces), "
     "Tiendanimal, Kiwoko, Leroy Merlin; webs propias de Dispensari Veterinari del Vallès, VetClínic, Centre Veterinari Les Torres, "
     "Centre Veterinari Rubí Nord, Petslovers, Agrifauna, Visa Garden."],
    ["Información publicada por el comercio", "El Comerç de Rubí (rubicomerc.cat, directorio de la asociación de comerciantes), "
     "páginas de Facebook de los comercios."],
    ["Directorios y otras fuentes", "Páginas Amarillas, QDQ, Cylex, Tiendeo, Caninoteca, Empresite, Iberinform, Botigues de Catalunya, "
     "Foursquare, Waze, Mapcarta (OpenStreetMap), prensa local y sectorial (Ràdio Rubí, Diari de Rubí, Rubí TV, Food Retail, Revista Aral)."],
    ["Fuentes NO accesibles", "Desde el entorno de trabajo no se pudo abrir ninguna página web directamente: la herramienta de descarga "
     "falló por DNS y el proxy bloqueó Google Maps, OpenStreetMap/Overpass/Nominatim, Páginas Amarillas y rubi.cat. La verificación web se "
     "hizo con un buscador que devuelve resúmenes de las páginas: los datos se han leído en esos resúmenes, no en la página original. "
     "Google Maps no se ha consultado. Alternativa usada: Overture Maps, descargado directamente de su bucket público de S3."],
    [""],
    ["2. ESTRATEGIA DE BÚSQUEDA"],
    ["Fase 1, descubrimiento", "(a) Extracción de Overture con un bbox amplio (13.489 lugares) y filtro geométrico dentro del polígono "
     "de Rubí (1.854). (b) Preselección por categoría (veterinarian, pet_store, grocery_store, supermarket, garden, health_food, "
     "farm, etc.) y por palabras clave en castellano y catalán (mascot, animal, pinso/pienso, gos, gat, agro, garden, jardí, herbol, "
     "natur, nombres de cadenas): 255 candidatos. (c) Revisión manual de los 255. (d) Búsquedas web complementarias en castellano y "
     "catalán ('tienda de animales Rubí', 'botiga animals Rubí pinso', 'veterinaris Rubí', 'herbolari Rubí', 'pinsos Rubí', 'vivers "
     "Rubí', cada cadena + Rubí, rubicomerc mascotes). Se encontraron 9 establecimientos que no estaban en Overture (Rubí Fauna, Cuida "
     "tus mascotas, Animal Paradis, El Flamenco Rosa, Gosset.net, Gaju Ricart, Consultorio Can Oriol, Herboristeria Majó, Condis Joan "
     "Maragall)."],
    ["Fase 2, consolidación", "Fusión de duplicados por nombre, dirección, teléfono y URL (14 fusiones). Comprobación de pertenencia "
     "municipal con el polígono y las direcciones; exclusión de establecimientos de Sant Cugat, Terrassa y Gavà."],
    ["Fase 3, verificación", "Búsqueda individual de cada establecimiento de mascotas, veterinario, garden y cadena: existencia, "
     "dirección, actividad, contacto y venta de alimentación."],
    ["Fase 4, clasificación", "Normalización a los 31 campos V1 con los valores permitidos; validación automática en el script "
     "scripts/build_rubi_v1.py."],
    ["Fase 5, control de calidad", "Validación de dominios de valores, de la coherencia entre campos (p. ej. perros='si' exige "
     "vende_alimentacion='si'), de unicidad de nombre+dirección y de coordenadas, y de coordenadas dentro del rango del municipio."],
    [""],
    ["3. FECHA DE LA INVESTIGACIÓN", "2026-10-14 (recopilación y verificación web). Cartografía: Overture 2026-09-23.0."],
    [""],
    ["4. CRITERIOS DE VERIFICACIÓN"],
    ["confirmado", "Identidad y ubicación verificadas en al menos dos fuentes sin contradicciones, actividad reciente, y evidencia "
     "específica de venta de alimentación procedente de la web oficial o de información publicada por el propio comercio."],
    ["parcialmente_confirmado", "Identidad y ubicación verificadas, pero falta un dato clave (normalmente la venta de alimentación o el "
     "detalle por especie) o hay discrepancias menores."],
    ["no_verificado", "Una sola fuente, fuentes antiguas o existencia actual dudosa."],
    ["Alimentación", "'si' solo con evidencia explícita para ese establecimiento. A los supermercados de cadena se les asigna "
     "'no_verificado': el surtido corporativo es un indicio, no una confirmación por sucursal. Los veterinarios no se presuponen vendedores."],
    ["Estado", "'activo' con evidencia reciente (web oficial o localizador actual, ficha cartográfica 2026 con alta confianza y "
     "directorios coherentes). 'cerrado' solo con confirmación (web oficial). En el resto, 'pendiente_revision'."],
    ["fecha_ultima_verificacion", "2026-10-14 solo en los registros contrastados individualmente por búsqueda web ese día; NULL en los "
     "que solo proceden de Overture."],
    ["Coordenadas", "Del punto de Overture del propio establecimiento (no geocodificación de direcciones). NULL cuando no hay punto o "
     "cuando dos puntos se contradicen (ALDI). No se ha hecho comprobación visual sobre ortofoto."],
    ["metodo_verificacion", "'fuente oficial' cuando la web o el localizador de la empresa confirman el dato; 'fuente externa' cuando "
     "son directorios o cartografía; NULL cuando no hubo verificación individual."],
    ["Teléfonos", "Formato E.164. Los números corporativos o de atención al cliente de una cadena no se asignan al establecimiento."],
    [""],
    ["5. LIMITACIONES DETECTADAS"],
    ["L1", "No se pudieron abrir páginas web ni Google Maps: las comprobaciones dependen de resúmenes del buscador. Puede haber errores "
     "de lectura en horarios o números."],
    ["L2", "No ha habido contacto telefónico ni visita: ningún dato se ha confirmado con el comercio ('comercio' y 'manual' no se usan)."],
    ["L3", "La venta de alimentación en supermercados no está verificada por sucursal."],
    ["L4", "Varios comercios pequeños solo aparecen en directorios antiguos (rubicomerc tiene fichas de más de 10 años); su actividad actual "
     "es incierta."],
    ["L5", "Overture contiene errores de clasificación (peluquerías como veterinarios), direcciones mal asignadas y duplicados; se "
     "corrigieron los detectados, pero puede quedar alguno."],
    ["L6", "La coordenada de Overture puede proceder de la geocodificación de la fuente original (Meta/Microsoft), no de una "
     "observación en campo."],
    [""],
    ["6. POSIBLES SESGOS O PROBLEMAS DE COBERTURA"],
    ["Cobertura", "No se puede demostrar una cobertura completa. Es probable que falten tiendas de barrio y bazares sin presencia digital, "
     "gasolineras con tienda, ferreterías o droguerías con sección de mascotas, y tiendas de alimentación de barrio. Los polígonos "
     "industriales (Can Jardí, Cova Solera, Rubí Sud) pueden tener distribuidores de piensos sin venta al público identificada."],
    ["Sesgo digital", "Los comercios con más presencia online (cadenas, veterinarios con web) están mejor verificados que los "
     "independientes pequeños."],
    ["Frontera municipal", "El área comercial de Can Sant Joan (Tiendanimal, Leroy Merlin, Caprabo) se percibe como 'Rubí' y aparece "
     "como tal en algunos directorios, pero pertenece a Sant Cugat."],
    [""],
    ["7. RECOMENDACIONES PARA OTROS MUNICIPIOS"],
    ["R1", "Empezar por la cartografía (Overture u OSM) con el polígono municipal oficial (ICGC/IGN) para delimitar el término "
     "municipal, y después enriquecer con localizadores oficiales de las cadenas."],
    ["R2", "Usar el directorio local de comercio (equivalente a rubicomerc) y el censo de actividades económicas o de licencias municipales "
     "si el ayuntamiento lo publica."],
    ["R3", "Hacer una ronda telefónica breve con un guion fijo (¿vende pienso de perro?, ¿de gato?, ¿otros?, ¿horario?) para pasar los "
     "registros a 'comercio', y una visita para supermercados representativos de cada cadena."],
    ["R4", "Disponer de acceso a Google Places API o similar para estado de actividad y coordenadas, y documentar el place_id."],
    ["R5", "Mantener la hoja de descartados y duplicados para no reincorporar errores, y reutilizar scripts/build_rubi_v1.py como plantilla."],
    ["R6", "Habilitar pronto la reclamación de la ficha por el comercio (reclamado_por_comercio) como vía de verificación continua."],
]


if __name__ == "__main__":
    print(main(sys.argv[1] if len(sys.argv) > 1 else "."))
